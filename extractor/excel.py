"""Build the Excel workbook from output/json/*.json."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .validate import SUBJ_CODE

REVIEW_BELOW = 0.8

COLUMNS = [
    ("Q_ID", 30, lambda p, q: q["q_id"]),
    ("Exam", 10, lambda p, q: p["exam"]),
    ("Year", 7, lambda p, q: p["year"]),
    ("Date", 11, lambda p, q: p["date"]),
    ("Shift", 9, lambda p, q: p["shift"]),
    ("Subject", 12, lambda p, q: q["subject"]),
    ("Section", 8, lambda p, q: q["section"]),
    ("Q_No", 6, lambda p, q: q["q_no"]),
    ("Chapter", 28, lambda p, q: q["chapter"]),
    ("Topic", 26, lambda p, q: q["topic"]),
    ("Question_Type", 13, lambda p, q: q["question_type"]),
    ("Question_Style", 18, lambda p, q: q["question_style"]),
    ("Question", 70, lambda p, q: q["question_text"]),
    ("Option_A", 25, lambda p, q: _opt(q, 0)),
    ("Option_B", 25, lambda p, q: _opt(q, 1)),
    ("Option_C", 25, lambda p, q: _opt(q, 2)),
    ("Option_D", 25, lambda p, q: _opt(q, 3)),
    ("Extra_Options", 15, lambda p, q: " | ".join(q["options"][4:])),
    ("Correct_Answer", 10, lambda p, q: q["answer"]),
    ("Answer_Source", 18, lambda p, q: q.get("answer_source", "")),
    ("Explanation", 90, lambda p, q: q.get("explanation", "")),
    ("Explanation_Check", 12, lambda p, q: q.get("explanation_check", "") or ("missing" if not q.get("explanation") else "")),
    ("Has_Diagram", 8, lambda p, q: "Yes" if q["has_diagram"] else "No"),
    ("Diagram_File", 30, lambda p, q: "; ".join(q.get("diagram_files", []))),
    ("Has_Equation", 8, lambda p, q: "Yes" if "$" in q["question_text"] + "".join(q["options"]) else "No"),
    ("Difficulty_Est", 10, lambda p, q: q["difficulty"]),
    ("Confidence", 10, lambda p, q: q["confidence"]),
    ("Verified_Twice", 9, lambda p, q: "Yes" if q.get("verified") else ""),
    ("Issues", 40, lambda p, q: "; ".join(q["issues"])),
    ("Source_File", 40, lambda p, q: Path(p["file"]).name),
    ("Page_No", 7, lambda p, q: q["page_index"] + 1),
]


def _opt(q, i):
    return q["options"][i] if i < len(q["options"]) else ""


HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
LOW_FILL = PatternFill("solid", fgColor="FCE4D6")


def _sheet(ws, rows, link_col):
    ws.append([c[0] for c in COLUMNS])
    for cell in ws[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for p, q in rows:
        ws.append([c[2](p, q) for c in COLUMNS])
        r = ws.max_row
        if q.get("diagram_files"):
            cell = ws.cell(r, link_col)
            cell.hyperlink = q["diagram_files"][0]
            cell.font = Font(color="0563C1", underline="single")
        if q["confidence"] < REVIEW_BELOW:
            ws.cell(r, [c[0] for c in COLUMNS].index("Confidence") + 1).fill = LOW_FILL
    for i, (_, w, _) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def build(out_dir: Path, xlsx: Path):
    papers = [json.loads(f.read_text()) for f in sorted((out_dir / "json").glob("*.json"))]
    papers.sort(key=lambda d: (d["paper"]["date"] or "", d["paper"]["shift"]))
    rows = [(d["paper"], q) for d in papers for q in d["questions"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    link_col = [c[0] for c in COLUMNS].index("Diagram_File") + 1
    _sheet(ws, rows, link_col)
    _sheet(wb.create_sheet("Needs_Review"), [r for r in rows if r[1]["confidence"] < REVIEW_BELOW], link_col)

    s = wb.create_sheet("Papers_Summary")
    head = ["Source_File", "Exam", "Date", "Shift", "Pages", "Text_Layer", "Total_Q"] + \
           [f"{c}_Q" for c in SUBJ_CODE.values()] + ["With_Diagram", "Explained", "Answer_Missing", "Needs_Review", "Avg_Confidence", "Paper_Issues", "Model"]
    s.append(head)
    for cell in s[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for d in papers:
        p, qs = d["paper"], d["questions"]
        s.append([Path(p["file"]).name, p["exam"], p["date"], p["shift"], p["pages"], "Yes" if p["text_layer"] else "No (scanned)",
                  len(qs)] + [sum(q["subject"] == subj for q in qs) for subj in SUBJ_CODE] +
                 [sum(q["has_diagram"] for q in qs), sum(bool(q.get("explanation")) for q in qs), sum(not q["answer"] for q in qs),
                  sum(q["confidence"] < REVIEW_BELOW for q in qs),
                  round(sum(q["confidence"] for q in qs) / max(len(qs), 1), 3), "; ".join(d["issues"]), p["model"]])
    for i, w in enumerate([60, 10, 11, 9, 7, 12, 8, 7, 7, 7, 12, 10, 14, 12, 14, 60, 18], 1):
        s.column_dimensions[get_column_letter(i)].width = w
    s.freeze_panes = "B2"

    ch = wb.create_sheet("Chapter_Counts")
    ch.append(["Subject", "Chapter", "Questions", "With_Diagram"])
    for cell in ch[1]:
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    counts = {}
    for _, q in rows:
        k = (q["subject"], q["chapter"])
        c = counts.setdefault(k, [0, 0])
        c[0] += 1
        c[1] += q["has_diagram"]
    for (subj, chap), (n, dg) in sorted(counts.items()):
        ch.append([subj, chap, n, dg])
    for i, w in enumerate([14, 45, 11, 13], 1):
        ch.column_dimensions[get_column_letter(i)].width = w

    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)
    return len(rows), sum(r[1]["confidence"] < REVIEW_BELOW for r in rows)
