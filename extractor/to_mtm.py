"""Convert output/JEE_Questions.xlsx into MockTestMitra's Supabase shape.

Writes to output/mocktestmitra/:
  MTM_JEE_Main_2021.xlsx   Tests + Questions tabs, headers = DB column names, in the
                           conventions of MockTestMitra's scripts/google-sheets-sync.gs
                           (position 1-based, correct as A-D letters, TRUE/FALSE).
  tests.json, questions.json
                           DB-ready rows (position 0-based, correct 0-3 / numeric) for a
                           direct PostgREST upsert (see upsert_mtm.py).
  figures/                 diagram PNGs keyed exactly as they go into the
                           question-assets Storage bucket.

Run:  python3 -m extractor.to_mtm
"""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "output" / "JEE_Questions.xlsx"
OUT = ROOT / "output" / "mocktestmitra"

SUPABASE_URL = "https://ipsogtxbeigczuhdhqkq.supabase.co"
BUCKET = "question-assets"
ASSET_PREFIX = "jee-main/pyq-2021"

EXAM_SLUG = "jee-main"
EXAM_VALUE = "JEE Main"          # questions.exam — must equal practice-test-generator examValue
SUBJ_PREFIX = {"Physics": "phy", "Chemistry": "chem", "Mathematics": "math"}
SHIFT_NO = {"Morning": 1, "Evening": 2}
MONTHS = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
TOTAL_TIME = 3600                # 30 questions per subject; JEE allots 60 min per subject

TEST_COLS = ["id", "exam_slug", "title", "subject", "topic", "total_time", "is_active", "visibility"]
QUESTION_COLS = ["test_id", "position", "type", "question_text", "option_a", "option_b",
                 "option_c", "option_d", "correct", "tolerance", "explanation", "diagram",
                 "difficulty", "topic", "subject", "exam", "source", "is_pyq", "year",
                 "math_format"]

DISPLAY_RE = re.compile(r"\$\$(.+?)\$\$", re.S)
INLINE_RE = re.compile(r"(?<![\\$])\$(?!\$)(.+?)(?<![\\$])\$(?!\$)", re.S)
ANSWER_LINE_RE = re.compile(r"\n?Answer:[^\n]*\s*$")


def latexify(s):
    """$$…$$ stays display math; inline $…$ becomes \\(…\\) (single $ is disabled on MTM)."""
    if s is None:
        return None
    s = str(s)
    held = []

    def hold(m):
        held.append(m.group(0))
        return f"\x00{len(held) - 1}\x00"

    s = DISPLAY_RE.sub(hold, s)
    s = INLINE_RE.sub(lambda m: r"\(" + m.group(1) + r"\)", s)
    if "$" in s:
        raise ValueError(f"unbalanced $ left after conversion: {s[:120]!r}")
    return re.sub(r"\x00(\d+)\x00", lambda m: held[int(m.group(1))], s)


def test_id(r):
    y, mo, d = r["Date"].split("-")
    return f"jee-{SUBJ_PREFIX[r['Subject']]}-pyq-{y}-{mo}-{d}-s{SHIFT_NO[r['Shift']]}"


def test_title(r):
    y, mo, d = r["Date"].split("-")
    return (f"JEE Main {y} PYQ – {int(d)} {MONTHS[int(mo)]} Shift {SHIFT_NO[r['Shift']]}"
            f" – {r['Subject']}")


def explanation(r, is_num):
    body = ANSWER_LINE_RE.sub("", str(r["Explanation"]).rstrip())
    ans = str(r["Correct_Answer"]).strip()
    head = f"Correct answer: {ans}" if is_num else f"Correct answer: ({ans.lower()})"
    return latexify(head + "\n" + body)


def number(v):
    f = float(str(v).strip())
    return int(f) if f.is_integer() else f


def build():
    ws = openpyxl.load_workbook(SRC, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    rows = [dict(zip(hdr, v)) for v in it if v[0]]

    tests, questions, figures, skipped = {}, [], [], []
    next_pos = {}
    for r in rows:
        if str(r["Correct_Answer"]).strip().lower() == "bonus":
            skipped.append(r["Q_ID"])            # no valid key: correct is NOT NULL, so leave out
            continue
        tid = test_id(r)
        if tid not in tests:
            tests[tid] = {"id": tid, "exam_slug": EXAM_SLUG, "title": test_title(r),
                          "subject": r["Subject"], "topic": "Previous Year Paper",
                          "total_time": TOTAL_TIME, "is_active": True, "visibility": "public"}
        is_num = r["Question_Type"] == "Numerical"
        pos = next_pos.get(tid, 0)                  # contiguous 0-based, source order A1..B10
        next_pos[tid] = pos + 1
        qtype = ("numerical" if is_num else
                 "assertion_reason" if r["Question_Style"] == "Assertion-Reason" else "standard")
        if is_num:
            correct, letter = number(r["Correct_Answer"]), None
        else:
            letter = str(r["Correct_Answer"]).strip().upper()
            correct = "ABCD".index(letter)
        diagram = None
        if r.get("Diagram_File"):
            rel = str(r["Diagram_File"]).split(";")[0].strip()
            key = f"{ASSET_PREFIX}/{Path(rel).parent.name}/{Path(rel).name}"
            figures.append((ROOT / "output" / rel, key))
            diagram = f"{SUPABASE_URL}/storage/v1/object/public/{BUCKET}/{key}"
        questions.append({
            "test_id": tid, "position": pos, "type": qtype,
            "question_text": latexify(r["Question"]),
            **{f"option_{c}": (None if is_num else latexify(r[f"Option_{c.upper()}"]))
               for c in "abcd"},
            "correct": correct, "tolerance": 0 if is_num else None,
            "explanation": explanation(r, is_num), "diagram": diagram,
            "difficulty": (r["Difficulty_Est"] or "").lower() or None,
            "topic": r["Chapter"], "subject": r["Subject"], "exam": EXAM_VALUE,
            "source": f"JEE Main {r['Year']}", "is_pyq": True, "year": int(r["Year"]),
            "math_format": "latex", "_letter": letter,
        })

    check(tests, questions, figures)
    if skipped:
        print("skipped (Bonus / dropped, no valid key):", ", ".join(skipped))
    return list(tests.values()), questions, figures


def check(tests, questions, figures):
    seen = set()
    for q in questions:
        k = (q["test_id"], q["position"])
        assert k not in seen, f"duplicate position {k}"
        seen.add(k)
        for f in ("question_text", "option_a", "option_b", "option_c", "option_d", "explanation"):
            v = q[f] or ""
            assert v.count(r"\(") == v.count(r"\)"), (k, f, "unbalanced \\( \\)")
            assert v.count("$$") % 2 == 0, (k, f, "odd $$")
        assert q["explanation"].lower().startswith("correct answer:"), k
        if q["type"] != "numerical":
            assert all(q[f"option_{c}"] for c in "abcd"), (k, "missing option")
    per = {}
    for q in questions:
        per.setdefault(q["test_id"], []).append(q["position"])
    for t, p in per.items():
        assert sorted(p) == list(range(len(p))), (t, "positions not contiguous")
    for src, _ in figures:
        assert src.exists(), f"missing figure {src}"


def write_sheet(tests, questions):
    wb = openpyxl.Workbook()
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")

    def tab(ws, cols, rows):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = bold, fill
        for row in rows:
            ws.append([row[c] for c in cols])
        ws.freeze_panes = "A2"

    t = wb.active
    t.title = "Tests"
    tab(t, TEST_COLS, tests)
    sheet_q = []
    for q in questions:                       # sheet conventions of google-sheets-sync.gs
        s = dict(q)
        s["position"] = q["position"] + 1
        s["correct"] = q["_letter"] if q["_letter"] else q["correct"]
        sheet_q.append(s)
    tab(wb.create_sheet("Questions"), QUESTION_COLS, sheet_q)
    wb.save(OUT / "MTM_JEE_Main_2021.xlsx")


def main():
    tests, questions, figures = build()
    OUT.mkdir(parents=True, exist_ok=True)
    write_sheet(tests, questions)
    db_q = [{k: v for k, v in q.items() if not k.startswith("_")} for q in questions]
    (OUT / "tests.json").write_text(json.dumps(tests, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT / "questions.json").write_text(json.dumps(db_q, ensure_ascii=False, indent=1), encoding="utf-8")
    fig_dir = OUT / "figures"
    if fig_dir.exists():
        shutil.rmtree(fig_dir)
    for src, key in figures:
        dst = fig_dir / key
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
    print(f"tests {len(tests)} | questions {len(questions)} | figures {len(figures)} -> {OUT}")


if __name__ == "__main__":
    main()
